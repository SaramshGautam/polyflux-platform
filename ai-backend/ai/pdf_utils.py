import io
import fitz  # PyMuPDF
from pptx import Presentation
from docx import Document

def extract_text_from_pdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    full_text = ""
    for page in doc:
        full_text += page.get_text()
    return full_text.strip()

def chunk_text(text: str, chunk_size: int = 400, overlap: int = 50) -> list[str]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i : i + chunk_size])
        chunks.append(chunk)
        i += chunk_size - overlap
    return chunks


def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extract all visible text from a .pptx slide deck: slide text boxes,
    tables, and speaker notes (slide by slide, in order)."""
    prs = Presentation(io.BytesIO(file_bytes))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append(notes)
    return "\n".join(parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all paragraph and table text from a .docx document."""
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()
